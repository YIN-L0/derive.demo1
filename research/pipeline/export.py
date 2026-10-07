#!/usr/bin/env python3
"""Build review and import exports from the canonical research database.

    python3 research/pipeline/export.py

Writes exports/ (Dérive-schema CSV, extended CSV, Excel review workbook, summary.json)
and verified/ (the subset whose status is `verified`). Reads only research/.
The flat exports merge relationship types into one cell for display, as the
existing workbook does; the canonical records keep them separate.
"""
import collections, csv, json
from pathlib import Path

R = Path(__file__).resolve().parent.parent
J = lambda p, k: json.loads((R / p).read_text(encoding="utf-8"))[k]

works = {w["work_id"]: w for w in J("normalized/works.json", "works")}
creators = {c["creator_id"]: c for c in J("normalized/creators.json", "creators")}
places = {p["place_id"]: p for p in J("normalized/places.json", "places")}
rels = J("normalized/relationships.json", "relationships")
claims = {c["claim_id"]: c for c in J("normalized/claims.json", "claims")}
evidence = {e["evidence_id"]: e for e in J("normalized/evidence.json", "evidence")}
routes = J("normalized/routes.json", "routes")
sources = {s["source_id"]: s for s in J("sources/source_registry.json", "sources")}
conflicts = J("conflicts/conflicts.json", "conflicts")
HEADER = json.loads((R / "schemas/derive_atlas_v3_header.json").read_text(encoding="utf-8"))["columns"]
crosswalk = list(csv.DictReader(open(R / "normalized/crosswalk_derive_sites.csv", encoding="utf-8")))
review = list(csv.DictReader(open(R / "conflicts/review_queue.csv", encoding="utf-8")))

REL_LABEL = {"textual_place": "Textual place", "inferred_place": "Inferred place", "narrative_location": "Narrative location",
             "biographical_place": "Biographical place", "inspiration_place": "Inspiration place", "composition_place": "Composition place",
             "filming_location": "Filming location", "adaptation_location": "Adaptation location"}
REL_ORDER = list(REL_LABEL)
TIER = {1: "A", 2: "B", 3: "C", 4: "D"}
EV_LABEL = {"primary_text_quote": "primary text (quotation checked against the source file)", "author_statement": "author's own statement",
            "structured_record": "structured dataset record", "inscription": "plaque inscription", "secondary_statement": "secondary statement"}
STATUS_ORDER = ["verified", "strongly_supported", "probable", "unresolved", "conflicting"]


def work_label(wid):
    w = works[wid]
    return f"{w['title']} ({w['year']})"


# which works touch which place (new works + the existing Dérive atlas)
place_works = collections.defaultdict(set)
for r in rels:
    for wid in ([r["work_id"]] if r["work_id"] else r["related_work_ids"]):
        place_works[r["place_id"]].add(wid)
existing_by_place = {c["place_id"]: c for c in crosswalk}
route_node = {}
for rt in routes:
    for s in rt["stops"]:
        route_node[(rt["work_id"], s["place_id"])] = f"{s['order']} — {rt['name']} ({rt['locator']})"


def rows_for_export():
    groups = collections.OrderedDict()
    for r in rels:
        for wid in ([r["work_id"]] if r["work_id"] else r["related_work_ids"]):
            groups.setdefault((wid, r["place_id"]), []).append(r)
    out = []
    for (wid, pid), rs in groups.items():
        rs = sorted(rs, key=lambda r: REL_ORDER.index(r["relationship_type"]))
        p = places[pid]
        text_rel = next((r for r in rs if r["relationship_type"] in ("textual_place", "inferred_place", "narrative_location")), None)
        cl = [claims[c] for r in rs for c in dict.fromkeys(r["claim_ids"])]
        cl = list({c["claim_id"]: c for c in cl}.values())
        ev = [evidence[e] for c in cl for e in c["evidence_ids"]]
        g = p["geometry"]
        coords = "" if g["lat"] is None else (("" if g["precision"] == "exact" else "Approx. ") + f"{g['lat']:.5f}, {g['lon']:.5f}" + ("" if g["precision"] == "exact" else f" ({g['precision']}-level point)"))
        name = p["display_name"]
        niw = text_rel["name_in_work"] if text_rel else None
        if niw and niw.lower().replace("the ", "").replace("le ", "").replace("la ", "") not in name.lower():
            name = f"{name} — in the work: “{niw}”"
        hist = [n["name"] for n in p["names"] if n["kind"] == "historic"]
        then_now = []
        if hist:
            then_now.append(f"Formerly {', '.join(hist)}; now {p['canonical_name']}.")
        if p["existence_status"] in ("defunct", "demolished"):
            then_now.append(f"Status: {p['existence_status']}.")
        if p["existence_status"] == "unknown":
            then_now.append("Present-day identity or survival not established.")
        bio = [c["statement"] for c in cl if c["claim_type"] in ("biographical", "composition", "commemoration")]
        ident = [c["statement"] for c in cl if c["claim_type"] == "identification"]
        others = sorted(work_label(w) for w in place_works[pid] if w != wid)
        if pid in existing_by_place:
            x = existing_by_place[pid]
            others += [f"{t} [existing Dérive site {x['derive_site_id']}; {x['match_type']}]" for t in x["existing_derive_works"].split("; ")]
        kinds = list(dict.fromkeys(EV_LABEL[e["evidence_type"]] for e in ev if e["stance"] == "supports"))
        tiers = sorted({sources[e["source_id"]]["tier"] for e in ev if e["stance"] == "supports"})
        statuses = list(dict.fromkeys(f"{REL_LABEL[r['relationship_type']].lower()}: {r['verification_status']}" for r in rs))
        if len({r["verification_status"] for r in rs}) == 1:
            statuses = [rs[0]["verification_status"]]
        urls = list(dict.fromkeys(sources[e["source_id"]]["url"] for e in ev if sources[e["source_id"]]["url"].startswith("http")))
        locs = list(dict.fromkeys(l for r in rs for l in r["locators"]))
        row = {
            "Corpus / Work": work_label(wid), "Place": name, "Address": p["address_current"] or "", "Coordinates": coords,
            "Place Relationship Type": "; ".join(dict.fromkeys(REL_LABEL[r["relationship_type"]] for r in rs)),
            "Mention Type": (text_rel or {}).get("mention_type") or "", "Timestamp / Chapter": "; ".join(locs[:8]) + (" …" if len(locs) > 8 else ""),
            "Characters / People": next((r["characters_people"] for r in rs if r["characters_people"]), ""),
            "Textual / Dialogue Anchor": f"“{text_rel['textual_anchor']}”" if text_rel and text_rel["textual_anchor"] else "",
            "What Happens Here": next((r["what_happens"] for r in rs if r["what_happens"]), ""),
            "Narrative Consequence": "", "Spatial Function": (text_rel or {}).get("spatial_role") or "", "Why This Place Matters": "",
            "Then & Now": " ".join(then_now), "Practical Information": "", "Biographical / Historical Layer": " ".join(bio),
            "Inspiration / Adaptation Layer": " ".join(ident), "Route Order / Story Node": route_node.get((wid, pid), ""),
            "Nearby / Related Works": "; ".join(others), "Evidence Chain": " + ".join(kinds), "Source Tier": "/".join(TIER[t] for t in tiers),
            "Confidence": "; ".join(statuses), "Sources": " ; ".join(urls),
        }
        ext = {"work_id": wid, "place_id": pid, "wpr_ids": "; ".join(r["wpr_id"] for r in rs), "claim_ids": "; ".join(c["claim_id"] for c in cl),
               "wikidata": p["external_ids"].get("wikidata") or "", "coord_precision": g["precision"] or "", "coord_source": g["coord_source_id"] or "",
               "derive_site_id": p["derive_site_id"] or "", "review_flags": "; ".join(dict.fromkeys(f for r in rs for f in r["review_flags"]))}
        out.append((row, ext))
    return out


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=header)
        w.writeheader()
        w.writerows(rows)


exp = rows_for_export()
write_csv(R / "exports/derive_atlas_export.csv", HEADER, [r for r, _ in exp])
ext_header = ["work_id", "place_id", "wpr_ids"] + HEADER + ["claim_ids", "wikidata", "coord_precision", "coord_source", "derive_site_id", "review_flags"]
write_csv(R / "exports/derive_atlas_export_extended.csv", ext_header, [{**r, **e} for r, e in exp])

# verified subset
flat = lambda d: {k: ("; ".join(map(str, v)) if isinstance(v, list) else v) for k, v in d.items() if not isinstance(v, dict) and k != "locators_structured"}
vr = [flat(r) for r in rels if r["verification_status"] == "verified"]
vc = [flat(c) for c in claims.values() if c["verification_status"] == "verified"]
write_csv(R / "verified/relationships_verified.csv", list(vr[0]), vr)
write_csv(R / "verified/claims_verified.csv", list(vc[0]), vc)

# summary
summary = {"works": {}, "shared_places": [], "totals": {}}
for wid, w in works.items():
    mine = [r for r in rels if wid in ([r["work_id"]] if r["work_id"] else r["related_work_ids"])]
    by_place = collections.defaultdict(list)
    for r in mine:
        by_place[r["place_id"]].append(r["verification_status"])
    # a place counts once, at the weakest status among its relationships to this work
    place_status = collections.Counter(max(v, key=STATUS_ORDER.index) for v in by_place.values())
    summary["works"][w["title"]] = {
        "places": len(by_place), "relationships": len(mine),
        "places_by_weakest_status": {s: place_status.get(s, 0) for s in STATUS_ORDER},
        "relationships_by_type": dict(collections.Counter(r["relationship_type"] for r in mine)),
        "relationships_by_status": dict(collections.Counter(r["verification_status"] for r in mine)),
        "primary_text_coverage": w["primary_text_coverage"],
    }
for pid, ws in sorted(place_works.items(), key=lambda kv: (-len(kv[1]), kv[0])):
    ex = existing_by_place.get(pid)
    n = len(ws) + (len(ex["existing_derive_works"].split("; ")) if ex and ex["match_type"] == "same" else 0)
    if n >= 2:
        summary["shared_places"].append({"place_id": pid, "name": places[pid]["canonical_name"], "new_works": sorted(works[w]["title"] for w in ws),
                                         "existing_derive_works": ex["existing_derive_works"].split("; ") if ex and ex["match_type"] == "same" else [],
                                         "existing_derive_site_id": ex["derive_site_id"] if ex and ex["match_type"] == "same" else None})
summary["totals"] = {"places": len(places), "places_with_coordinates": sum(p["geometry"]["lat"] is not None for p in places.values()),
                     "relationships": len(rels), "claims": len(claims), "claims_by_status": dict(collections.Counter(c["verification_status"] for c in claims.values())),
                     "evidence": len(evidence), "primary_quotes_machine_checked": sum(e["machine_checked_against_source"] for e in evidence.values()),
                     "sources": len(sources), "sources_used_as_evidence": sum(s["used_as_evidence"] for s in sources.values()),
                     "routes": len(routes), "conflicts": len(conflicts), "review_queue": len(review)}
(R / "exports/summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

# Excel review workbook (for people; not canonical)
try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None
if openpyxl:
    wb = openpyxl.Workbook()

    def sheet(title, header, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for r in rows:
            ws.append([("; ".join(map(str, v)) if isinstance(v, list) else json.dumps(v, ensure_ascii=False) if isinstance(v, dict) else v) for v in (r.get(h) for h in header)])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2F3B4C")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, min(48, max(12, len(h) + 4)))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        return ws

    ws = wb.active
    ws.title = "README"
    for line in ["Dérive cultural-place research — review workbook", "", "Generated from research/normalized/*.json by research/pipeline/export.py.",
                 "This workbook is for reading and commenting. It is NOT the database: edits here are lost on the next export.",
                 "Change the JSON (or tell the researcher what to change), then re-run validate.py and export.py.", "",
                 "Status words (verified, strongly_supported, probable, unresolved, conflicting) are defined in schemas/vocabularies.json.",
                 "'verified' means the evidence standard was met. It does not mean a person has reviewed the row.", "",
                 "Sheets: Dérive export (same 23 columns as the product workbook) · Places · Relationships · Claims · Evidence · Routes · Sources · Conflicts · Review queue · Shared places"]:
        ws.append([line])
    ws.column_dimensions["A"].width = 120
    ws["A1"].font = Font(bold=True, size=14)
    wide = {"Textual / Dialogue Anchor": 50, "What Happens Here": 45, "Sources": 60, "Biographical / Historical Layer": 50, "Inspiration / Adaptation Layer": 50, "statement": 70, "quote": 70, "notes": 50, "issue": 70, "suggested_action": 60, "contains": 60, "current_assessment": 70}
    sheet("Dérive export", HEADER, [r for r, _ in exp], wide)
    sheet("Places", ["place_id", "canonical_name", "place_type", "existence_status", "names", "address_current", "lat", "lon", "coord_precision", "coord_source_id", "wikidata", "part_of", "derive_site_id", "derive_match_type", "notes"],
          [{**p, "names": "; ".join(f"{n['name']} [{n['kind']}]" for n in p["names"][1:]), "lat": p["geometry"]["lat"], "lon": p["geometry"]["lon"], "coord_precision": p["geometry"]["precision"],
            "coord_source_id": p["geometry"]["coord_source_id"], "wikidata": p["external_ids"].get("wikidata")} for p in places.values()], wide)
    sheet("Relationships", ["wpr_id", "work_id", "creator_id", "place_id", "relationship_type", "verification_status", "spatial_role", "mention_type", "name_in_work", "locators", "characters_people", "textual_anchor", "what_happens", "period", "identification_basis", "claim_ids", "review_flags"], rels, {**wide, "textual_anchor": 50, "what_happens": 45, "identification_basis": 50})
    sheet("Claims", ["claim_id", "claim_type", "verification_status", "statement", "work_id", "place_id", "creator_id", "best_source_tier", "evidence_ids", "conflict_id", "notes", "assessed_by", "assessed_date"], list(claims.values()), wide)
    sheet("Evidence", ["evidence_id", "claim_id", "source_id", "evidence_type", "stance", "locator", "record_id", "quote", "quote_lang", "quote_is_verbatim", "machine_checked_against_source", "accessed", "note"], list(evidence.values()), {**wide, "note": 50})
    sheet("Routes", ["route_id", "work_id", "route_name", "locator", "order", "place_id", "name_in_work", "note", "route_status"],
          [{"route_id": rt["route_id"], "work_id": rt["work_id"], "route_name": rt["name"], "locator": rt["locator"], "route_status": rt["verification_status"], **s} for rt in routes for s in rt["stops"]])
    sheet("Sources", ["source_id", "tier", "title", "institution_or_author", "url", "source_type", "structured", "formats", "license_reuse", "works_covered", "contains", "usefulness", "automatable", "access_status", "used_as_evidence", "notes"], list(sources.values()), {**wide, "title": 50, "license_reuse": 50, "usefulness": 50, "automatable": 45})
    sheet("Conflicts", ["conflict_id", "title", "status", "claim_ids", "positions", "current_assessment", "open_questions"],
          [{**c, "positions": " || ".join(f"{p['position']} [{', '.join(p['sources'])}]" for p in c["positions"])} for c in conflicts], {**wide, "title": 50, "positions": 70, "open_questions": 60})
    sheet("Review queue", list(review[0]), review, {**wide, "subject": 40})
    sheet("Shared places", ["place_id", "name", "new_works", "existing_derive_works", "existing_derive_site_id"], summary["shared_places"], {"name": 40, "new_works": 60, "existing_derive_works": 50})
    wb.save(R / "exports/derive_research_review.xlsx")

print(json.dumps({k: {"places": v["places"], **v["places_by_weakest_status"]} for k, v in summary["works"].items()}, indent=1))
print("shared places:", len(summary["shared_places"]), "| export rows:", len(exp), "| xlsx:", bool(openpyxl))
